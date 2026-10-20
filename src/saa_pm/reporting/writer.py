"""Write study results: one CSV per table, charts, an Excel workbook with an
index sheet, and a short plain-language SUMMARY.md."""

from __future__ import annotations

import os
import re
from typing import Dict, List

import pandas as pd

from .._logging import logger


class ResultWriter:
    def __init__(self, out_dir: str):
        self.out = out_dir
        os.makedirs(out_dir, exist_ok=True)
        self.tables: Dict[str, pd.DataFrame] = {}

    def _path(self, rel: str) -> str:
        p = os.path.join(self.out, rel)
        os.makedirs(os.path.dirname(p), exist_ok=True)
        return p

    def table(self, rel: str, df) -> None:
        if df is None or (hasattr(df, "empty") and df.empty):
            return
        if isinstance(df, pd.Series):
            df = df.to_frame()
        self.tables[rel] = df
        df.to_csv(self._path(rel + ".csv"), float_format="%.6g")

    def text(self, rel: str, txt: str) -> None:
        with open(self._path(rel), "w", encoding="utf-8") as f:
            f.write(txt + "\n")

    def chart_path(self, name: str) -> str:
        return self._path(os.path.join("charts", name))

    def workbook(self, name: str = "saa_results.xlsx") -> str | None:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            logger.warning(
                "openpyxl is not installed - Excel workbook skipped (pip install 'saa-pm[report]')"
            )
            return None
        path = self._path(name)
        index_rows, used = [], set()
        with pd.ExcelWriter(path, engine="openpyxl") as xl:
            for rel, df in self.tables.items():
                base = re.sub(r"^\d+_", "", rel).replace("/", "_")[:31]
                sheet, k = base, 1
                while sheet.lower() in used:
                    suffix = f"~{k}"
                    sheet, k = base[: 31 - len(suffix)] + suffix, k + 1
                used.add(sheet.lower())
                out = df.copy()
                for c in out.columns:  # Excel cannot store tz-aware datetimes
                    if hasattr(out[c], "dt") and getattr(out[c].dt, "tz", None) is not None:
                        out[c] = out[c].dt.tz_localize(None)
                if len(out) > 5000:  # monthly long tables: keep CSV only
                    index_rows.append(
                        dict(sheet="(csv only)", source=rel + ".csv", rows=len(out), cols=out.shape[1])
                    )
                    continue
                out.to_excel(xl, sheet_name=sheet)
                index_rows.append(dict(sheet=sheet, source=rel + ".csv", rows=len(out), cols=out.shape[1]))
            pd.DataFrame(index_rows).to_excel(xl, sheet_name="INDEX", index=False)
            wb = xl.book
            wb.move_sheet("INDEX", offset=-(len(wb.sheetnames) - 1))
        return path

    def summary(self, res) -> str:
        d = res.data
        L: List[str] = ["# SAA study - summary", ""]
        L.append(
            f"Data as of **{d.asof:%Y-%m-%d}**, {d.universe.n} asset classes, "
            f"portfolio value **{d.holdings.total:,.0f}**, "
            f"{len(d.common_returns)} months of history."
        )
        L.append("")
        if res.allocation:
            a = res.allocation
            cfg = a.state.optimizer
            L += [
                "## Goals (in priority order)",
                "",
                f"1. Expected return >= **{cfg.target_return:.2%}** (CPI {a.state.goals.cpi:.2%} "
                f"+ {a.state.goals.return_spread:.2%})",
                f"2. Expected income >= **{cfg.target_income:.2%}**",
                f"3. Volatility between **{cfg.vol_lo:.1%}** and **{cfg.vol_hi:.1%}**",
                "",
            ]
            s = a.summary()
            L += [
                "## Allocation",
                "",
                "| portfolio | exp. return | income | volatility | illiquid |",
                "|---|---:|---:|---:|---:|",
            ]
            for k, r in s.iterrows():
                L.append(f"| {k} | {r.exp_return:.2%} | {r.income:.2%} | {r.vol:.2%} | {r.illiquid:.1%} |")
            L += [
                "",
                f"- Strategic target: {a.strategic.solution.selection_rule}",
                f"- Today: {a.today.solution.selection_rule}",
            ]
            if len(a.ladder_table):
                L.append(
                    f"- Funding ladder pre-funds {a.ladder.total:.1%} of portfolio value "
                    f"for scheduled capital calls"
                )
            for n in a.notes:
                L.append(f"- note: {n}")
            L.append("")
        if res.backtest_metrics is not None:
            m = res.backtest_metrics
            L += [
                "## Walk-forward backtest",
                "",
                "| run | ann. return | ann. vol | max drawdown | income |",
                "|---|---:|---:|---:|---:|",
            ]
            for k in m.columns:
                L.append(
                    f"| {k} | {m.loc['ann_return', k]:.2%} | {m.loc['ann_vol', k]:.2%} | "
                    f"{m.loc['max_drawdown', k]:.1%} | {m.loc['realised_income', k]:.2%} |"
                )
            L.append("")
        T = res.tables
        if "private_markets/performance_total" in T:
            p = T["private_markets/performance_total"].iloc[0]
            L += [
                "## Private markets programme",
                "",
                f"Projected IRR {p.irr:.2%}, TVPI {p.tvpi:.2f}x, DPI {p.dpi:.2f}x on "
                f"{p.paid_in:,.0f} paid in.",
                "",
            ]
        if "wealth/probabilities" in T:
            pr = T["wealth/probabilities"]
            last = pr.index.max()
            col = "p_beat_target" if "p_beat_target" in pr else pr.columns[0]
            L += [
                "## Wealth projection",
                "",
                f"Probability of compounding at or above the return target over {last} years: "
                f"**{pr.loc[last, col]:.0%}**; probability of a real loss: "
                f"{pr.loc[last, 'p_real_loss']:.0%}.",
                "",
            ]
        if res.validation.warnings:
            L += ["## Data warnings", ""] + [f"- {w}" for w in res.validation.warnings] + [""]
        L += [
            "## Files",
            "",
            "- `saa_results.xlsx` - every table, one sheet each (see INDEX)",
            "- `charts/` - figures",
            "- numbered folders - the same tables as CSV",
            "",
        ]
        txt = "\n".join(L)
        self.text("SUMMARY.md", txt)
        return txt
